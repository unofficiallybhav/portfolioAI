import os
from pathlib import Path
from dotenv import load_dotenv
from pydantic import BaseModel
from groq import Groq
from pypdf import PdfReader
from docx import Document
import json
import re
import time

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("API key not found")
client=Groq(api_key=my_api_key)
model="llama-3.3-70b-versatile"

def llm_call(sys,user):
    msgs=[{"role":"system","content":sys},{"role":"user","content":user}]
    response=client.chat.completions.create(model=model,messages=msgs,response_format={"type":"json_object"})
    return response.choices[0].message.content

# Extract text from pdf file
def extract_pdf(path):
    reader = PdfReader(path)
    full_text = ""
    for page in reader.pages:
        text = page.extract_text()
        if text:
            full_text += text + "\n"
    return full_text

# Extract text from docx file
def extract_word(path):
    reader=Document(path)
    full_text=""
    for para in reader.paragraphs:
        if para.text.strip():
            full_text+=para.text+"\n"
    
    for table in reader.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    full_text+=cell.text+"\n"
    return full_text

# Read resume file
def read_file(path):
    if ".pdf" in path:
        return extract_pdf(path)
    elif ".docx" in path:
        return extract_word(path)
    return None

# Create candidate profile
class profile(BaseModel):
    name:str
    skills:list
    project_details:list[str]
    cgpa:float
    education:str
    certifications:str
    work_exp:str
    internship_details:str
profile_schema=profile.model_json_schema()

def extract_profile(resume_path=r"resume.pdf"):
    print("CREATING CANDIDATE PROFILE\n")
    sys=f"""
        You are a experienced talent agent/headhunter.
        Your job is to curate a profile for a candidate using their resume.
        Only stick to the information mentioned in the resume and dont invent anything that is not true.
        NO EXTRA EXPLANATION SHOULD BE PRESENT.
        Include details about all the projects done in project_details field and internship details in internship_details field.
        Return a JSON output strictly matching the schema {profile_schema}.
    """
    user=f"""
        I am Bhavyy Khurana and this is my resume {(read_file(resume_path))}
    """
    raw_profile=llm_call(sys,user)
    profile_data=json.loads(raw_profile)
    candidate=profile(**profile_data)
    return candidate

# Extract JD from text/file
class JD(BaseModel):
    role:str
    skills:list
    education_qualifications:str
    work_exp_qualifications:str
JD_schema=JD.model_json_schema()

def extract_JD(jd):
    print("\nEXTRACTING JOB DESCRIPTION\n")
    system_prompt=f"""
        You are a HR assistant.
        Your job is to parse the job description for a role and extract information regarding the skills required.
        Return a JSON output strictly matching the schema {JD_schema} and no extra text. 
        Do not invent any extra skills that are not mentioned.
    """
    user_prompt=f"""
        Go through {jd} and report your findings.
    """
    raw_jd=llm_call(system_prompt,user_prompt)
    jd=json.loads(raw_jd)
    return JD(**jd)

# Match skills of candidate with job description
class Match(BaseModel):
    score:int
    verdict:str
match_schema=Match.model_json_schema()

def match_skills(role,candidate):
    print("CREATING VERDICT")
    system_prompt=f"""
        You are a HR assistant.
        Your job is to match the skills of a candidate with the skills required for a job, give score between 1 and 100 and short verdict explaining the score.
        Return a JSON output strictly matching the schema {match_schema} and no extra text.
    """
    user_prompt=f"""
        The skills required for the job role are {role} and the skills of the candidate are {candidate}. Report your findings.
    """
    raw_score=llm_call(system_prompt,user_prompt)
    score=json.loads(raw_score)
    return Match(**score)


ACTION_RE = re.compile(r"Action:\s*(\w+)")

class AgentState(BaseModel):
    candidate:profile
    jd:JD|None=None
    match_result:Match|None=None
    prompt:str
    last_obs:str|None=None
    history:list[dict]=[]

    def build_prompt(self)->str:
        prompt=f"""
            You are a portfolio chatbot.
            Your only job is to answer any questions related to the technical background
            of Bhavyy Khurana using {self.candidate} as a knowledge base. 
            If prompted to give any information which is not part of the knowledge base
            then mention that you are just a portfolio chatbot and cant help with this query.
            You will judged on how honestly you can answer the questions. You should not
            invent any new information or make up any lies regarding the information presented to you.

            The tools at your disposal are:
            extract_JD
                which can parse an job description provided by the user
            match_skills
                which can compare the candidate's skills with that neede for any job description
            
            WORKFLOW:
            1. Decide what you need to do next.
            2. Only call one tool at a time.
            3. After writing an ACTION, STOP immediately.
            4. Never guess or invent a tool result.
            5. Wait until you recive an observation.
            6. Then make your next action.
            7. When task is complete give final answer.

            Never write
            extract_JD(jd)
            match_skills(candidate,role)

            When you need a tool, use EXACTLY this format:

            Thought: ...
            Action: extract_JD

            or

            Thought: ...
            Action: match_skills

            After receiving an Observation, continue reasoning until the task is complete.

            When you are ready to answer the user, output ONLY:

            Final Answer:
            <the complete response to the user>

            Do not write any explanation before "Final Answer:".
            Do not summarize after "Final Answer:".
            The entire user-visible answer must appear after "Final Answer:".
            
            Only follow the FORMAT when a tool call is needed if information is already 
            present then just respond.

            Current Agent State:
            Candidate profile is present
        """
        if self.jd:
            prompt+="""Job description is present"""
        if self.match_result:
            prompt+="""Match result is present"""
        return prompt

    def build_tools(memory: "AgentState"):    
        def run_extract_JD():
            memory.jd = extract_JD(memory.prompt)
            return memory.jd.model_dump_json()
    
        def run_match_skills():
            if memory.jd is None:
                return "Error: no job description parsed yet. Call extract_JD first."
            memory.match_result = match_skills(memory.jd, memory.candidate)
            return memory.match_result.model_dump_json()
    
        return {"extract_JD": run_extract_JD, "match_skills": run_match_skills}

def run_agent(memory: AgentState):
    tools = memory.build_tools()
    msgs = [{"role": "system", "content": memory.build_prompt()}]
    msgs += memory.history
    msgs.append({"role": "user", "content": memory.prompt})

    for step in range(5):
        reply = client.chat.completions.create(
            model=model, messages=msgs, temperature=0
        ).choices[0].message.content
        print(reply)
        msgs.append({"role": "assistant", "content": reply})

        if "Final Answer" in reply:
            memory.history = msgs[1:][-20:]
            return reply

        found = ACTION_RE.search(reply)
        if not found:
            observation = ("No valid Action found. Reply with 'Action: tool_name' "
                        "or 'Final Answer: ...'.")
        else:
            tool_name = found.group(1)
            if tool_name not in tools:
                observation = f"Unknown tool {tool_name!r}. Available: {list(tools)}"
            else:
                try:
                    observation = tools[tool_name]()
                except Exception as exc:
                    observation = f"Tool {tool_name} failed: {exc}"

        print(f"\nObservation: {observation}")
        memory.last_obs = observation
        msgs.append({"role": "user", "content": "Observation: " + str(observation)})
        msgs[0]["content"] = memory.build_prompt()

    memory.history = msgs[1:][-20:]
    return "Stopped: maximum steps reached without a final answer."

def main():
    memory = AgentState(candidate=extract_profile(), prompt="")
    print(f"Portfolio bot for {memory.candidate.name}. Ctrl-C to quit.\n")
    while True:
        try:
            user_input = input("> ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if not user_input:
            continue
        memory.prompt = user_input
        print(run_agent(memory))

if __name__=="__main__":
    main()